"""Export the rule registry and deadline table to docs/legal/ (xlsx and md).

The registry is the source of truth once a rule is coded; this keeps the table
Benjamin reviews in lockstep with the code. Run: python tools/export_rule_table.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from dueproc.rules import registry

OUT = Path(__file__).resolve().parent.parent / "docs" / "legal"
COLUMNS = [
    "ID",
    "Layer",
    "Severity",
    "Citation",
    "Quoted authority",
    "Plain-language flag",
    "Suggested action",
    "URL",
    "Review status",
]


# Appendix B, with the unit of every clock confirmed (or marked open) in Week 1.
DEADLINES = [
    (
        "notice",
        "Written notice of suspension with reasons and review right",
        "Suspension decision",
        "Immediately; flagged if not received by the next school day",
        "school days (1)",
        "105 ILCS 5/10-22.6(b)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "max_suspension",
        "Maximum suspension per incident",
        "First day of removal",
        "10",
        "school days",
        "105 ILCS 5/10-22.6(b)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "support_services",
        "Support services determination",
        "Suspension decision",
        "Applies when OSS exceeds 4",
        "school days",
        "105 ILCS 5/10-22.6(b-25)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "interventions",
        "Documented exhaustion of interventions",
        "Suspension decision",
        "Applies when OSS exceeds 3",
        "school days",
        "105 ILCS 5/10-22.6(b-20)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "expulsion_max",
        "Maximum expulsion",
        "Expulsion decision",
        "2",
        "calendar years",
        "105 ILCS 5/10-22.6(a)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "records",
        "Records access after parent request",
        "Receipt of written request",
        "15",
        "school days",
        "105 ILCS 10/5 (ISSRA)",
        "OPEN: brief says 15 school days; verify current ISSRA text (may be business days)",
    ),
    (
        "tenth_day",
        "IDEA services must continue",
        "Cumulative removals in the school year",
        "After 10",
        "school days",
        "34 CFR 300.530(d)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "change_of_placement",
        "Change of placement threshold",
        "Removals in the school year",
        "More than 10 consecutive, or a pattern totalling more than 10",
        "school days",
        "34 CFR 300.536",
        "per Appendix B; pattern is a team determination, tool treats it as unknown",
    ),
    (
        "safeguards_notice",
        "Parent notification and procedural safeguards notice",
        "Decision to make a removal that is a change of placement",
        "Same day",
        "same day",
        "34 CFR 300.530(h)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "mdr",
        "Manifestation determination review",
        "Decision to change placement",
        "10",
        "school days",
        "34 CFR 300.530(e)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "expedited_hearing",
        "Expedited due process hearing",
        "Parent or district request",
        "Hearing within 20; decision within 10 after hearing",
        "school days",
        "34 CFR 300.532(c)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "iaes_max",
        "Interim alternative educational setting, special circumstances",
        "Removal",
        "Not more than 45",
        "school days",
        "34 CFR 300.530(g)",
        "per Appendix B; re-verify at legal review",
    ),
    (
        "firearm",
        "Firearm expulsion",
        "Expulsion decision",
        "Not less than 1 year unless modified case by case",
        "calendar years",
        "105 ILCS 5/10-22.6(d)",
        "per Appendix B; re-verify at legal review",
    ),
]
DEADLINE_COLUMNS = [
    "Obligation ID",
    "Obligation",
    "Clock starts",
    "Length",
    "Unit",
    "Authority",
    "Unit status",
]


def rows() -> list[list[str]]:
    return [
        [
            r.id,
            r.layer.value,
            r.severity.value,
            r.citation,
            r.quote,
            r.flag,
            r.action,
            r.url,
            "reviewed" if r.reviewed else "pending legal review; quote to be re-verified",
        ]
        for r in registry()
    ]


def write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "rules_v1"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows():
        ws.append(row)
    widths = [8, 10, 12, 28, 60, 60, 50, 40, 22]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    wb.save(path)


def write_deadlines(xlsx: Path, md: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "deadlines_v1"
    ws.append(DEADLINE_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in DEADLINES:
        ws.append(list(row))
    for i, w in enumerate([18, 45, 35, 40, 15, 26, 50]):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    wb.save(xlsx)
    lines = [
        "# Deadline table v1",
        "",
        "Every clock has a start event, a length, a unit, and an authority. Generated by",
        "`tools/export_rule_table.py`; the timeline constants in `src/dueproc/timeline.py` must match.",
        "",
        "| " + " | ".join(DEADLINE_COLUMNS) + " |",
        "|" + "---|" * len(DEADLINE_COLUMNS),
    ]
    lines += ["| " + " | ".join(row) + " |" for row in DEADLINES]
    md.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_md(path: Path) -> None:
    lines = [
        "# Rule table v1",
        "",
        "Generated from the rule registry by `tools/export_rule_table.py`. Do not edit by hand.",
        "Every quote is pending re-verification against the current ilga.gov / eCFR text and",
        "pending the Gate 3 legal review.",
        "",
        "| ID | Severity | Citation | Flag | Action |",
        "|---|---|---|---|---|",
    ]
    for r in rows():
        lines.append(f"| {r[0]} | {r[2]} | {r[3]} | {r[5]} | {r[6]} |")
    lines += ["", "## Quoted authority", ""]
    for r in rows():
        lines += [f'**{r[0]}** ({r[3]}, [source]({r[7]})): "{r[4]}"', ""]
    path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    write_xlsx(OUT / "rules_v1.xlsx")
    write_md(OUT / "rules_v1.md")
    write_deadlines(OUT / "deadlines_v1.xlsx", OUT / "deadlines_v1.md")
    print(f"wrote {len(rows())} rules to {OUT}")
