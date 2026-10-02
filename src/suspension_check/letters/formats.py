"""Render a plain-text letter as a printable HTML page or a DOCX file.

The text template is the only source of wording; these functions only change
its presentation. Both outputs end with a separate "for you" page carrying the
disclaimer and the resource card, clearly marked as not part of the letter.
Nothing is written to disk: DOCX is built in memory and returned as bytes.
"""

from __future__ import annotations

import io
import re
from html import escape

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from suspension_check.report import DISCLAIMER, RESOURCES

__all__ = ["paragraphs", "to_docx", "to_html"]

_KEEP_NOTE = "Keep this page. It is for you, not part of the letter."


def paragraphs(text: str) -> list[list[str]]:
    """Split on blank lines; each paragraph keeps its own line breaks."""
    blocks = re.split(r"\n\s*\n", text.strip())
    return [b.split("\n") for b in blocks if b.strip()]


def to_html(text: str, title: str) -> str:
    body = "\n".join(
        "<p>" + "<br>".join(escape(line) for line in block) + "</p>" for block in paragraphs(text)
    )
    resources = "\n".join(
        f'<li><a href="{escape(r["url"])}">{escape(r["name"])}</a>: {escape(r["for"])}</li>'
        for r in RESOURCES
    )
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)}</title>
<style>
  body {{ font: 12pt/1.5 Georgia, "Times New Roman", serif; color: #111; background: #fff;
    max-width: 7in; margin: 0.75in auto; padding: 0 16px; }}
  p {{ margin: 0 0 1em; }}
  .notes {{ border-top: 2px solid #999; margin-top: 2em; padding-top: 1em;
    font: 11pt/1.45 system-ui, sans-serif; }}
  .notes h2 {{ font-size: 13pt; margin: 0 0 .5em; }}
  @media print {{ .notes {{ break-before: page; border: 0; }} a {{ color: inherit; }} }}
</style>
</head>
<body>
<main class="letter">
{body}
</main>
<section class="notes" aria-label="Notes for you">
<h2>{escape(_KEEP_NOTE)}</h2>
<p>Fill in every part in [brackets] before you send the letter. Keep a copy and write
down the date you sent it and how (email, hand delivery, or mail).</p>
<p>{escape(DISCLAIMER)}</p>
<ul>
{resources}
</ul>
</section>
</body>
</html>
"""


def to_docx(text: str, title: str) -> bytes:
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = ""  # no identifying metadata in the file
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(12)
    for block in paragraphs(text):
        p = doc.add_paragraph()
        for i, line in enumerate(block):
            if i:
                p.add_run().add_break(WD_BREAK.LINE)
            p.add_run(line)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(_KEEP_NOTE, level=2)
    doc.add_paragraph(
        "Fill in every part in [brackets] before you send the letter. Keep a copy and "
        "write down the date you sent it and how."
    )
    doc.add_paragraph(DISCLAIMER)
    for r in RESOURCES:
        doc.add_paragraph(f"{r['name']}: {r['for']} ({r['url']})", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)  # persistence-ok: in-memory buffer, never a file
    return buf.getvalue()
