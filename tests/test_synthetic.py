"""The synthetic logs on disk must agree with the plans their expected metrics come from."""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

import pytest
from openpyxl import load_workbook
from synthetic.generate import LOG_A, LOG_B, LOG_C, SUBJECTIVE, expected_metrics

SYN = Path(__file__).parent.parent / "synthetic"


def _read(plan) -> list[dict[str, str]]:
    inv = {v: k for k, v in plan.headers.items()}
    path = SYN / f"{plan.name}.{plan.fmt}"
    if plan.fmt == "csv":
        with path.open(encoding="utf-8") as fh:
            return [{inv[k]: v for k, v in row.items()} for row in csv.DictReader(fh)]
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    return [{inv[h]: str(v) for h, v in zip(rows[0], r, strict=True)} for r in rows[1:]]


@pytest.mark.parametrize("plan", [LOG_A, LOG_B, LOG_C], ids=lambda p: p.name)
def test_log_matches_plan(plan):
    rows = [r for r in _read(plan) if r["action_type"] == "OSS"]
    per_student = Counter((r["race_ethnicity"], r["student_id"]) for r in rows)
    students = Counter(g for g, _ in per_student)
    repeats = Counter(g for (g, _), n in per_student.items() if n >= 2)
    subjective = defaultdict(int)
    for r in rows:
        subjective[r["race_ethnicity"]] += r["incident_category"] in SUBJECTIVE
    for g in plan.groups:
        assert students[g.group] == g.suspended_once + g.suspended_twice
        assert repeats[g.group] == g.suspended_twice
        incidents = g.suspended_once + 2 * g.suspended_twice
        assert subjective[g.group] == round(incidents * g.subjective_share)


@pytest.mark.parametrize("plan", [LOG_A, LOG_B, LOG_C], ids=lambda p: p.name)
def test_expected_json_is_current(plan):
    on_disk = json.loads((SYN / "expected" / f"{plan.name}.json").read_text())
    fresh = expected_metrics(plan)
    assert {k: on_disk[k] for k in fresh} == fresh


def test_small_cells_are_suppressed():
    m = expected_metrics(LOG_C)["by_race_ethnicity"]
    assert m["White"]["enrolled"] == "<10"
    assert all(v["students_suspended_per_100"] is None for v in m.values())
