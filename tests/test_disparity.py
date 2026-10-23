from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import Workbook, load_workbook

from suspension_check import disparity
from suspension_check.disparity import (
    OTHER,
    SUPPRESS_BELOW,
    PatternError,
    analyze,
    propose_mapping,
)

SYN = Path(__file__).parent.parent / "synthetic"
LOGS = [("school_a", "csv"), ("school_b", "xlsx"), ("school_c_small", "csv")]


def _expected(name):
    return json.loads((SYN / "expected" / f"{name}.json").read_text(encoding="utf-8"))


def _run(name, ext, **kw):
    exp = _expected(name)
    raw = (SYN / f"{name}.{ext}").read_bytes()
    return exp, analyze(raw, f"{name}.{ext}", {"race_ethnicity": exp["enrollment"]}, **kw)


def _csv(rows):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


# --- exact match against hand-derived expectations -------------------------------


@pytest.mark.parametrize(("name", "ext"), LOGS)
def test_metrics_match_expected_exactly(name, ext):
    exp, got = _run(name, ext)
    assert got["oss_incidents"] == exp["oss_incidents"]
    assert got["iss_rows_ignored"] == exp["iss_rows_ignored"]
    assert got["by_race_ethnicity"] == exp["by_race_ethnicity"]


@pytest.mark.parametrize(("name", "ext"), LOGS)
def test_headers_map_with_no_correction(name, ext):
    _, got = _run(name, ext)
    assert {"incident_date", "action_type", "days_assigned", "race_ethnicity"} <= set(
        got["fields_found"]
    )
    assert not got["warnings"]


# --- privacy -------------------------------------------------------------------------


@pytest.mark.parametrize(("name", "ext"), LOGS)
def test_names_and_ids_never_appear_in_output(name, ext):
    raw = (SYN / f"{name}.{ext}").read_bytes()
    if ext == "csv":
        rows = list(csv.reader(io.StringIO(raw.decode())))
    else:
        rows = [list(r) for r in load_workbook(io.BytesIO(raw)).active.iter_rows(values_only=True)]
    headers, body = rows[0], rows[1:]
    mapping = propose_mapping([str(h) for h in headers])
    secret_cols = [i for i, h in enumerate(headers) if mapping[str(h)] in ("name", "student_id")]
    assert secret_cols, "fixture should contain a name or ID column"
    secrets_ = {str(r[i]) for r in body for i in secret_cols if r[i]}
    _, got = _run(name, ext)
    out = json.dumps(got)
    leaked = [s for s in secrets_ if s in out]
    assert not leaked


def test_name_columns_are_dropped_and_reported_by_header_only():
    _, got = _run("school_a", "csv")
    assert got["mapping"]["Student Name"] == "name"
    assert "Student Name" in got["dropped_columns"]


def test_free_text_labels_cannot_surface_as_groups():
    rows = [["Date", "Action", "Days", "Race", "ID"]]
    rows += [["2025-10-01", "OSS", "2", "Black", f"s{i}"] for i in range(12)]
    rows += [["2025-10-02", "OSS", "1", "Jamal Carter", "x1"]]  # a name typed into Race
    got = analyze(_csv(rows), "log.csv", {"race_ethnicity": {"Black": 100, "White": 100}})
    assert "Jamal Carter" not in json.dumps(got)
    assert OTHER in got["by_race_ethnicity"]


def test_salt_is_fresh_per_call(monkeypatch):
    salts = []
    real = disparity.secrets.token_bytes
    monkeypatch.setattr(
        disparity.secrets, "token_bytes", lambda n: salts.append(real(n)) or salts[-1]
    )
    _run("school_c_small", "csv")
    _run("school_c_small", "csv")
    assert len(salts) == 2 and salts[0] != salts[1]


# --- suppression ---------------------------------------------------------------------


def test_small_school_suppresses_every_rate():
    _, got = _run("school_c_small", "csv")
    for m in got["by_race_ethnicity"].values():
        assert m["students_suspended"] == "<10"
        assert m["students_suspended_per_100"] is None
        assert m["risk_ratio"] is None
        assert m["composition_index"] is None
    assert all(v == "<10" for v in got["monthly"].values())


groups = st.sampled_from(["Black", "White", "Hispanic", "Asian"])


@settings(max_examples=60, deadline=None)
@given(
    st.lists(st.tuples(groups, st.integers(0, 30), st.integers(1, 6)), min_size=1, max_size=120),
    st.dictionaries(groups, st.integers(1, 400), min_size=1),
)
def test_no_cell_under_threshold_is_ever_shown(incidents, enrollment):
    rows = [["Date", "Action", "Days", "Race", "Student ID", "Behavior"]]
    for k, (g, sid, days) in enumerate(incidents):
        rows.append([f"2025-1{k % 3}-0{1 + k % 9}", "OSS", str(days), g, f"{g}{sid}", "Defiance"])
    got = analyze(_csv(rows), "x.csv", {"race_ethnicity": enrollment})
    for m in got["by_race_ethnicity"].values():
        for key in ("students_suspended", "enrolled"):
            v = m[key]
            assert v is None or v == "<10" or v >= SUPPRESS_BELOW
        if m["students_suspended"] == "<10" or m["enrolled"] in ("<10", None):
            assert m["students_suspended_per_100"] is None
            assert m["days_lost_per_100"] is None
            assert m["composition_index"] is None
            assert m["risk_ratio"] is None
    for v in [*got["monthly"].values(), *got["grade_bands"].values()]:
        assert v == "<10" or v >= SUPPRESS_BELOW


# --- mapping and validation ----------------------------------------------------------


def test_one_correction_rescues_an_unknown_header():
    rows = [["When", "Punishment", "Days", "Race"]]
    rows += [["2025-10-01", "OSS", "3", "Black"]] * 12
    raw = _csv(rows)
    with pytest.raises(PatternError, match="incident_date"):
        analyze(raw, "x.csv")
    with pytest.raises(PatternError, match="action_type"):
        analyze(raw, "x.csv", mapping={"When": "incident_date"})
    got = analyze(raw, "x.csv", mapping={"When": "incident_date", "Punishment": "action_type"})
    assert got["oss_incidents"] == 12


def test_unrecognizable_file_fails_clearly():
    raw = _csv([["Item", "Price"], ["apple", "1.00"], ["pear", "2.00"]])
    with pytest.raises(PatternError, match="needs a column"):
        analyze(raw, "groceries.csv")


def test_bad_rows_fail_unless_rare():
    good = [["2025-10-01", "OSS", "2", "Black", f"s{i}"] for i in range(40)]
    header = [["Date", "Action", "Days", "Race", "ID"]]
    one_bad = header + good + [["not a date", "OSS", "2", "Black", "z"]]
    got = analyze(_csv(one_bad), "x.csv")
    assert "1 rows were skipped" in got["warnings"][0] and "row" in got["warnings"][0]
    many_bad = header + good + [["nope", "OSS", "two", "Black", "z"]] * 10
    with pytest.raises(PatternError, match="10 of 50 rows"):
        analyze(_csv(many_bad), "x.csv")


def test_error_messages_never_echo_cell_values():
    rows = [["Date", "Action", "Days"]] + [["Jamal Carter", "OSS", "2"]] * 5
    with pytest.raises(PatternError) as exc:
        analyze(_csv(rows), "x.csv")
    assert "Jamal" not in str(exc.value)


def test_mapping_errors():
    raw = _csv([["Date", "Action", "Days"], ["2025-10-01", "OSS", "1"]])
    with pytest.raises(PatternError, match="not in the file"):
        analyze(raw, "x.csv", mapping={"Nope": "grade"})
    with pytest.raises(PatternError, match="not a Pattern Check field"):
        analyze(raw, "x.csv", mapping={"Date": "favorite_color"})
    with pytest.raises(PatternError, match="Two columns"):
        analyze(raw, "x.csv", mapping={"Action": "incident_date"})


def test_no_oss_rows():
    raw = _csv([["Date", "Action", "Days"], ["2025-10-01", "ISS", "1"]])
    with pytest.raises(PatternError, match="No out-of-school"):
        analyze(raw, "x.csv")


def test_without_student_ids_incidents_stand_in_and_repeat_is_unknown():
    rows = [["Date", "Action", "Days", "Race"]] + [["2025-10-01", "OSS", "2", "Black"]] * 12
    got = analyze(_csv(rows), "x.csv", {"race_ethnicity": {"Black": 100, "White": 100}})
    assert any("No student ID" in w for w in got["warnings"])
    assert got["by_race_ethnicity"]["Black"]["students_suspended"] == 12
    assert got["by_race_ethnicity"]["Black"]["repeat_share"] is None


def test_semicolon_csv_and_cp1252_and_us_dates():
    text = "Date;Action;Days;Race\n" + "10/01/2025;Out-of-school suspension;2;Hispanic\n" * 11
    got = analyze(text.encode("cp1252") + b"\n", "x.csv")
    assert got["oss_incidents"] == 11
    assert "Hispanic" in got["by_race_ethnicity"]


# --- limits and XLSX safety ----------------------------------------------------------


def test_size_and_row_limits(monkeypatch):
    with pytest.raises(PatternError, match="larger than"):
        analyze(b"x" * (disparity.MAX_BYTES + 1), "x.csv")
    monkeypatch.setattr(disparity, "MAX_ROWS", 5)
    rows = [["Date", "Action", "Days"]] + [["2025-10-01", "OSS", "1"]] * 7
    with pytest.raises(PatternError, match="more than 5 rows"):
        analyze(_csv(rows), "x.csv")


def test_xlsx_formulas_are_not_evaluated_or_echoed():
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Action", "Days", "Race"])
    for _ in range(40):
        ws.append(["2025-10-01", "OSS", 2, "Black"])
    ws.append(["2025-10-01", "OSS", '=HYPERLINK("http://evil","x")', "White"])
    buf = io.BytesIO()
    wb.save(buf)  # persistence-ok: in-memory buffer
    got = analyze(buf.getvalue(), "x.xlsx")
    # The formula has no cached value, so its row is skipped, never computed.
    assert got["oss_incidents"] == 40
    assert "1 rows were skipped" in got["warnings"][0]
    assert "evil" not in json.dumps(got)


def test_not_a_workbook():
    with pytest.raises(PatternError, match="not a readable XLSX"):
        analyze(b"PK\x03\x04garbage", "x.xlsx")


def test_empty_file():
    with pytest.raises(PatternError, match="no header"):
        analyze(b"", "x.csv")
